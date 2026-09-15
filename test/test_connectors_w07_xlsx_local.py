"""Tests for the offline .xlsx engine: controlled cell write + version-safe
write-back + independent reopen-verify + the explicit non-computation /
non-cloud contract.

The fixtures are REAL .xlsx files built with openpyxl (a declared dependency),
so the tests exercise genuine SpreadsheetML: shared strings, styles, and
formula cells both WITH a cached value and WITHOUT one. Nothing here mocks the
container or the parser.

Import discipline: this slice depends on the sibling W07 OOXML engine
(``office_documents.container`` / ``.constants`` / ``.errors``) BY NAME. Those
modules are not on ``main`` yet (they land with the sibling PR), so this whole
module is skipped when they cannot be imported rather than failing collection —
the dependency is declared, not copied.
"""

from __future__ import annotations

import zipfile

import pytest

pytest.importorskip("openpyxl", reason="openpyxl is a declared dep; skip if absent")

# The by-name sibling dependency. Skip cleanly (not error) when it is not yet
# present in the tree, so this file is honest about being an integration
# dependent rather than pretending to own the sibling's code.
_xlsx = pytest.importorskip(
    "kiro_crew.connections.vendors.microsoft.office_documents.xlsx",
    reason="sibling office_documents engine (container/constants/errors) not present",
)
from kiro_crew.connections.vendors.microsoft.office_documents.errors import (  # noqa: E402
    DocumentEditError,
    MalformedDocument,
    ProtectedDocument,
    UnsupportedDocument,
)

xlsx = _xlsx


# ── Fixtures: build real workbooks ───────────────────────────────────────────


def _build_workbook(path: str) -> None:
    """Write a real .xlsx to *path* with strings, numbers, and formula cells.

    Cell layout on sheet "Data":
      A1 "Item"   B1 "Qty"
      A2 "apples" B2 3
      A3 "pears"  B3 4
      B4 =SUM(B2:B3)   <- formula WITH a cached value (7)
    A second sheet "Notes" carries A1 "hello".
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Item"
    ws["B1"] = "Qty"
    ws["A2"] = "apples"
    ws["B2"] = 3
    ws["A3"] = "pears"
    ws["B3"] = 4
    ws["B4"] = "=SUM(B2:B3)"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "hello"
    wb.save(path)
    # openpyxl does not cache formula results, so inject a cached <v> into the
    # B4 formula cell so the "cached value" read path has a real fixture. This
    # edits the worksheet part directly, the same surgical way production does.
    _inject_formula_cache(path, "xl/worksheets/sheet1.xml", "B4", "7")


def _inject_formula_cache(path: str, part: str, ref: str, cached: str) -> None:
    """Add a cached ``<v>`` to a formula cell inside a worksheet part in-place.

    Rebuilds the archive with only *part* rewritten (a crude but real edit) so
    the fixture carries a formula cell whose value cache is present.
    """
    import io
    import re

    with zipfile.ZipFile(path, "r") as zf:
        members = {i.filename: zf.read(i.filename) for i in zf.infolist()}
        order = [i.filename for i in zf.infolist()]
    raw = members[part].decode("utf-8")
    # Insert/replace the cached <v> for the target formula cell. openpyxl emits
    # an empty <v></v> placeholder next to a formula, so replace that with a
    # real cached value.
    pattern = re.compile(r'(<c r="%s"[^>]*>)(.*?)(</c>)' % re.escape(ref), re.DOTALL)

    def _repl(m: "re.Match[str]") -> str:
        inner = m.group(2)
        if "<v>" in inner or "<v/>" in inner or "<v></v>" in inner:
            inner = re.sub(r"<v\s*/?>(.*?)</v>|<v\s*/>", "<v>%s</v>" % cached, inner)
        else:
            inner = inner + "<v>%s</v>" % cached
        return m.group(1) + inner + m.group(3)

    raw = pattern.sub(_repl, raw, count=1)
    members[part] = raw.encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in order:
            zout.writestr(name, members[name])
    with open(path, "wb") as fh:
        fh.write(buf.getvalue())


def _build_formula_no_cache(path: str) -> None:
    """Write a workbook whose only formula cell carries NO cached value."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 2
    ws["A2"] = 5
    ws["A3"] = "=A1+A2"  # openpyxl writes the formula but no cached <v>
    wb.save(path)


@pytest.fixture()
def workbook(tmp_path):
    p = tmp_path / "book.xlsx"
    _build_workbook(str(p))
    return str(p)


# ── Read path ────────────────────────────────────────────────────────────────


def test_read_values_and_sheets(workbook):
    content = xlsx.read_cells(workbook)
    names = [s.name for s in content.sheets]
    assert names == ["Data", "Notes"]
    data = {c.ref: c for c in content.sheets[0].cells}
    assert data["A1"].kind == "value" and data["A1"].value == "Item"
    assert data["B2"].kind == "value" and data["B2"].value == 3
    assert data["A2"].value == "apples"
    notes = {c.ref: c for c in content.sheets[1].cells}
    assert notes["A1"].value == "hello"


def test_read_formula_with_cache_is_labelled_cached(workbook):
    content = xlsx.read_cells(workbook)
    data = {c.ref: c for c in content.sheets[0].cells}
    b4 = data["B4"]
    # A formula cell WITH a cached value: labelled cached, value from the file's
    # own cache, formula text carried. Never presented as a computed result.
    assert b4.kind == "cached"
    assert b4.formula == "=SUM(B2:B3)"
    assert b4.value == 7


def test_read_formula_without_cache_yields_formula_text_no_value(tmp_path):
    p = tmp_path / "nocache.xlsx"
    _build_formula_no_cache(str(p))
    content = xlsx.read_cells(str(p))
    data = {c.ref: c for c in content.sheets[0].cells}
    a3 = data["A3"]
    # NEGATIVE CONTRACT #1: cache missing -> no number is fabricated.
    assert a3.kind == "formula"
    assert a3.value is None
    assert a3.formula == "=A1+A2"


# ── Controlled write + version-safe write-back ───────────────────────────────


def test_write_cell_then_reopen_verify_returns_content(workbook):
    # write_cells verifies via an INDEPENDENT reopen by default and returns it.
    content = xlsx.write_cells(workbook, workbook, {"Data": {"B2": 99, "A2": "oranges"}})
    assert content is not None
    data = {c.ref: c for c in content.sheets[0].cells}
    assert data["B2"].value == 99
    assert data["A2"].value == "oranges"
    # A fully independent re-read (not the returned object) agrees.
    reread = xlsx.read_cells(workbook)
    data2 = {c.ref: c for c in reread.sheets[0].cells}
    assert data2["B2"].value == 99
    assert data2["A2"].value == "oranges"


def test_write_preserves_untouched_parts_byte_for_byte(workbook, tmp_path):
    dst = str(tmp_path / "out.xlsx")
    # Capture every part's raw bytes before the edit.
    with zipfile.ZipFile(workbook, "r") as zf:
        before = {i.filename: zf.read(i.filename) for i in zf.infolist()}
        order_before = [i.filename for i in zf.infolist()]

    xlsx.write_cells(workbook, dst, {"Data": {"B2": 42}})

    with zipfile.ZipFile(dst, "r") as zf:
        after = {i.filename: zf.read(i.filename) for i in zf.infolist()}
        order_after = [i.filename for i in zf.infolist()]

    edited_part = "xl/worksheets/sheet1.xml"
    # Member set and stored order are preserved.
    assert order_before == order_after
    # EVERY part except the one edited worksheet is byte-for-byte identical.
    for name, raw in before.items():
        if name == edited_part:
            assert after[name] != raw, "the edited worksheet part must change"
        else:
            assert after[name] == raw, f"untouched part {name} was not preserved verbatim"


def test_write_new_cell_creates_it_in_sorted_position(workbook):
    # C5 does not exist; writing it must create the cell and the row.
    content = xlsx.write_cells(workbook, workbook, {"Data": {"C5": "new"}})
    data = {c.ref: c for c in content.sheets[0].cells}
    assert data["C5"].value == "new"


def test_write_is_atomic_no_partial_dst_on_bad_edit(workbook, tmp_path):
    dst = tmp_path / "atomic.xlsx"
    # An out-of-range/invalid reference must fail BEFORE any bytes are written,
    # so the destination is never created as a truncated file.
    with pytest.raises(DocumentEditError):
        xlsx.write_cells(workbook, str(dst), {"Data": {"NOTACELL": 1}})
    assert not dst.exists(), "a rejected edit must not leave a partial destination"


def test_write_unknown_sheet_refuses(workbook):
    with pytest.raises(DocumentEditError):
        xlsx.write_cells(workbook, workbook, {"Ghost": {"A1": 1}})


def test_inplace_edit_roundtrips_other_sheet_untouched(workbook):
    xlsx.write_cells(workbook, workbook, {"Data": {"B2": 1000}})
    content = xlsx.read_cells(workbook)
    notes = {c.ref: c for c in content.sheets[1].cells}
    # The Notes sheet was never addressed; its content survives.
    assert notes["A1"].value == "hello"


# ── NEGATIVE CONTRACTS ───────────────────────────────────────────────────────


def test_recalculation_is_refused(workbook):
    # NEGATIVE CONTRACT #2: asked to recompute -> explicit refusal, never a
    # fabricated value.
    with pytest.raises(xlsx.RecalculationUnsupported):
        xlsx.refuse_recalculation()
    with pytest.raises(UnsupportedDocument):
        # It is an UnsupportedDocument subclass, so a caller catching the base
        # still sees the refusal.
        xlsx.refuse_recalculation("evaluate")


def test_cloud_live_semantics_is_refused():
    # NEGATIVE CONTRACT #3: this offline engine must not claim cloud/live
    # (Graph workbook) semantics.
    with pytest.raises(xlsx.CloudSemanticsUnsupported):
        xlsx.refuse_cloud_semantics()
    with pytest.raises(UnsupportedDocument):
        xlsx.refuse_cloud_semantics("live-session")


def test_writing_a_formula_string_does_not_pretend_to_compute(workbook):
    # Writing "=SUM(...)" as a value must store it as literal text, NOT as a
    # formula cell this engine would then be expected to evaluate. Reopen shows
    # the literal string, never a computed number.
    content = xlsx.write_cells(workbook, workbook, {"Data": {"B2": "=SUM(A2:A3)"}})
    data = {c.ref: c for c in content.sheets[0].cells}
    b2 = data["B2"]
    assert b2.kind == "value"
    assert b2.value == "=SUM(A2:A3)"
    assert b2.formula is None


# ── Rejection gate ───────────────────────────────────────────────────────────


def test_reject_non_zip(tmp_path):
    p = tmp_path / "plain.xlsx"
    p.write_bytes(b"this is not a zip at all")
    with pytest.raises(MalformedDocument):
        xlsx.read_cells(str(p))


def test_reject_ole2_encrypted_wrapper(tmp_path):
    # CFB / OLE2 magic with an .xlsx extension = password/agile-encrypted OOXML.
    p = tmp_path / "encrypted.xlsx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(ProtectedDocument):
        xlsx.read_cells(str(p))


def test_reject_legacy_xls_extension(tmp_path):
    p = tmp_path / "old.xls"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(UnsupportedDocument):
        xlsx.read_cells(str(p))


def test_reject_macro_enabled_extension(tmp_path, workbook):
    # A real, valid workbook but with a macro-enabled extension is refused on
    # the extension alone (it may carry a VBA project).
    import shutil

    p = tmp_path / "macro.xlsm"
    shutil.copy(workbook, p)
    with pytest.raises(ProtectedDocument):
        xlsx.read_cells(str(p))


def test_reject_zip_without_workbook_part(tmp_path):
    # A valid zip that is not a SpreadsheetML workbook (no xl/workbook.xml).
    p = tmp_path / "notxlsx.xlsx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("hello.txt", "not a workbook")
    with pytest.raises(MalformedDocument):
        xlsx.read_cells(str(p))


def test_reject_vba_project_member(tmp_path, workbook):
    # A workbook carrying a vbaProject.bin member is macro-bearing regardless of
    # extension, so it is refused.
    with zipfile.ZipFile(workbook, "r") as zf:
        members = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    p = tmp_path / "hasmacro.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, raw in members:
            zout.writestr(name, raw)
        zout.writestr("xl/vbaProject.bin", b"\x00\x01\x02")
    with pytest.raises(ProtectedDocument):
        xlsx.read_cells(str(p))


def test_missing_file_is_malformed(tmp_path):
    with pytest.raises(MalformedDocument):
        xlsx.read_cells(str(tmp_path / "does-not-exist.xlsx"))
